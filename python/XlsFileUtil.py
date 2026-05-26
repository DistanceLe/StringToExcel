#!/usr/bin/python
# -*- coding: UTF-8 -*-

from openpyxl import load_workbook

class XlsFileUtil:
    'xlsx file util'

    def __init__(self, filePath):
        self.filePath = filePath
        self.data = load_workbook(filePath, read_only=True, data_only=True)

    def getAllTables(self):
        return self.data.worksheets

    def getTableByIndex(self, index):
        if index >= 0 and index < len(self.data.worksheets):
            return XlsTableWrapper(self.data.worksheets[index])
        else:
            print("XlsFileUtil error -- getTable:index")

    def getTableByName(self, name):
        return XlsTableWrapper(self.data[name])


class XlsTableWrapper:
    """Wraps an openpyxl worksheet to provide xlrd-compatible .row_values() and .col_values()."""

    def __init__(self, ws):
        self.ws = ws
        self._rows_cache = None
        self._max_row = 0
        self._max_col = 0

    def _ensure_loaded(self):
        if self._rows_cache is not None:
            return
        self._rows_cache = []
        for row in self.ws.iter_rows():
            self._rows_cache.append([cell.value for cell in row])
        self._max_row = len(self._rows_cache)
        self._max_col = max((len(r) for r in self._rows_cache), default=0)

    @property
    def nrows(self):
        self._ensure_loaded()
        return self._max_row

    @property
    def ncols(self):
        self._ensure_loaded()
        return self._max_col

    def row_values(self, row_index):
        self._ensure_loaded()
        if row_index < self._max_row:
            row = self._rows_cache[row_index]
            # Pad with empty strings to match xlrd behavior
            return row + [''] * (self._max_col - len(row))
        return [''] * self._max_col

    def col_values(self, col_index):
        self._ensure_loaded()
        result = []
        for row in self._rows_cache:
            if col_index < len(row):
                result.append(row[col_index] if row[col_index] is not None else '')
            else:
                result.append('')
        return result
